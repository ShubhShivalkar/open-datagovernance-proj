from rest_framework.decorators import api_view
from rest_framework.generics import ListCreateAPIView, RetrieveAPIView, ListAPIView
from rest_framework.response import Response
from rest_framework import status as http_status
from django.shortcuts import get_object_or_404
from django.http import HttpResponse
import io
import threading

from .models import ErasureRequest, ErasureAuditLog, AccessRequest
from .serializers import (
    ErasureRequestSerializer, ErasureRequestDetailSerializer, ErasureAuditLogSerializer,
    AccessRequestSerializer, AccessRequestDetailSerializer,
)
from . import services

REQUIRED_PHRASE = (
    "I approve of the full erasure and understand this would delete "
    "all the data and cannot be recovered"
)


class ErasureRequestListCreateView(ListCreateAPIView):
    serializer_class = ErasureRequestSerializer
    pagination_class = None

    def get_queryset(self):
        return ErasureRequest.objects.select_related("datasource", "requested_by").all()

    def perform_create(self, serializer):
        serializer.save(requested_by=self.request.user)


class ErasureRequestDetailView(RetrieveAPIView):
    queryset = ErasureRequest.objects.select_related("datasource", "requested_by").all()
    serializer_class = ErasureRequestDetailSerializer


@api_view(["POST"])
def trigger_plan(request, pk):
    req = get_object_or_404(ErasureRequest, pk=pk)
    if req.status not in ("draft", "planned", "failed"):
        return Response({"detail": "Plan can only be triggered from draft, planned, or failed status."}, status=400)
    # Reset to draft so polling waits until the background thread completes.
    # The thread sets status→"planned" only after ErasurePlan is persisted,
    # preventing a race where the frontend sees "planned" before the plan exists.
    req.status = "draft"
    req.save(update_fields=["status"])
    t = threading.Thread(target=services.build_erasure_plan, args=(str(req.id),), daemon=True)
    t.start()
    return Response({"detail": "Plan computation started."})


@api_view(["POST"])
def approve_request(request, pk):
    req = get_object_or_404(ErasureRequest, pk=pk)
    if req.status != "planned":
        return Response({"detail": "Request must be in 'planned' status to approve."}, status=400)
    phrase = request.data.get("confirmation_phrase", "").strip()
    if phrase != REQUIRED_PHRASE:
        return Response({"detail": "Confirmation phrase does not match."}, status=400)
    from .models import ErasureApproval
    ErasureApproval.objects.filter(request=req).delete()
    ErasureApproval.objects.create(
        request=req,
        approved_by=request.user,
        confirmation_phrase=phrase,
    )
    req.status = "approved"
    req.save(update_fields=["status"])
    return Response({"detail": "Request approved."})


@api_view(["POST"])
def execute_request(request, pk):
    req = get_object_or_404(ErasureRequest, pk=pk)
    if req.status != "approved":
        return Response({"detail": "Request must be approved before execution."}, status=400)
    t = threading.Thread(target=services.execute_erasure, args=(str(req.id),), daemon=True)
    t.start()
    return Response({"detail": "Erasure execution started."})


@api_view(["POST"])
def cancel_request(request, pk):
    req = get_object_or_404(ErasureRequest, pk=pk)
    if req.status not in ("draft", "planned", "pending_approval"):
        return Response({"detail": "Cannot cancel a request in this state."}, status=400)
    req.status = "cancelled"
    req.save(update_fields=["status"])
    return Response({"detail": "Request cancelled."})


class AuditLogListView(ListAPIView):
    queryset = ErasureAuditLog.objects.all()
    serializer_class = ErasureAuditLogSerializer
    pagination_class = None


# ---------------------------------------------------------------------------
# Data Access Request (DSAR) views
# ---------------------------------------------------------------------------

class AccessRequestListCreateView(ListCreateAPIView):
    serializer_class = AccessRequestSerializer
    pagination_class = None

    def get_queryset(self):
        return AccessRequest.objects.select_related("datasource", "requested_by").all()

    def perform_create(self, serializer):
        serializer.save(requested_by=self.request.user)


class AccessRequestDetailView(RetrieveAPIView):
    queryset = AccessRequest.objects.select_related("datasource", "requested_by").all()
    serializer_class = AccessRequestDetailSerializer


@api_view(["POST"])
def trigger_discovery(request, pk):
    req = get_object_or_404(AccessRequest, pk=pk)
    if req.status not in ("draft", "discovered", "failed"):
        return Response({"detail": "Discovery can only be triggered from draft, discovered, or failed status."}, status=400)
    req.status = "draft"
    req.save(update_fields=["status"])
    t = threading.Thread(target=services.discover_data_access, args=(str(req.id),), daemon=True)
    t.start()
    return Response({"detail": "Data discovery started."})


@api_view(["GET"])
def get_access_report(request, pk):
    req = get_object_or_404(AccessRequest, pk=pk)
    if req.status not in ("discovered", "report_ready", "sent"):
        return Response({"detail": "Discovery must be complete before a report can be generated."}, status=400)
    try:
        report = services.generate_access_report(str(req.id))
        req.status = "report_ready"
        req.save(update_fields=["status"])
        return Response(report)
    except Exception as exc:
        return Response({"detail": str(exc)}, status=500)


@api_view(["POST"])
def send_access_report(request, pk):
    req = get_object_or_404(AccessRequest, pk=pk)
    if req.status not in ("discovered", "report_ready", "sent"):
        return Response({"detail": "Discovery must be complete before sending the report."}, status=400)
    extra_emails = request.data.get("additional_emails", [])
    try:
        services.send_access_report(str(req.id), extra_emails)
        return Response({"detail": "Report dispatched."})
    except Exception as exc:
        return Response({"detail": str(exc)}, status=500)


@api_view(["GET"])
def export_access_xlsx(request, pk):
    req = get_object_or_404(AccessRequest, pk=pk)
    if req.status not in ("discovered", "report_ready", "sent"):
        return Response({"detail": "Discovery must be complete before exporting."}, status=400)
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        return Response({"detail": "openpyxl is required for xlsx export. Install it with: pip install openpyxl"}, status=500)

    report = services.generate_access_report(str(req.id))
    wb = openpyxl.Workbook()

    # ── Sheet 1: Summary ──────────────────────────────────────────────────────
    ws1 = wb.active
    ws1.title = "Summary"
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1D4ED8")
    label_font  = Font(bold=True)

    ws1.column_dimensions["A"].width = 28
    ws1.column_dimensions["B"].width = 60

    def hrow(row, label, value=""):
        ws1.cell(row=row, column=1, value=label).font = label_font
        ws1.cell(row=row, column=2, value=str(value) if value is not None else "")

    ws1["A1"] = "Data Access Report"
    ws1["A1"].font = Font(bold=True, size=14)
    ws1.merge_cells("A1:B1")
    ws1.row_dimensions[1].height = 22

    hrow(3, "Request ID",       report["request_id"])
    hrow(4, "Generated At",     report["generated_at"])
    hrow(5, "Data Source",      report["datasource_name"])
    hrow(6, "Subject ID",       report["subject"]["id"])
    hrow(7, "Subject Email",    report["subject"]["email"])
    hrow(8, "Identifier Field", report["subject"]["identifier_field"])
    hrow(9, "Summary",          report["summary"])
    hrow(11, "Legal Basis",     report["legal_basis"])
    hrow(12, "Notes",           report["notes"])
    ws1.cell(row=9, column=2).alignment = Alignment(wrap_text=True)
    ws1.cell(row=11, column=2).alignment = Alignment(wrap_text=True)
    ws1.row_dimensions[9].height = 36
    ws1.row_dimensions[11].height = 48

    # ── Sheet 2: Data Inventory ───────────────────────────────────────────────
    ws2 = wb.create_sheet("Data Inventory")
    ws2.column_dimensions["A"].width = 20
    ws2.column_dimensions["B"].width = 40
    ws2.column_dimensions["C"].width = 14
    ws2.column_dimensions["D"].width = 30
    ws2.column_dimensions["E"].width = 30
    ws2.column_dimensions["F"].width = 40

    col_headers = ["Table", "Business Purpose", "Record Count", "PII Columns", "Data Categories", "Relationship Path"]
    for col_i, header in enumerate(col_headers, start=1):
        cell = ws2.cell(row=1, column=col_i, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for row_i, entry in enumerate(report["data_inventory"], start=2):
        ws2.cell(row=row_i, column=1, value=entry["table"])
        ws2.cell(row=row_i, column=2, value=entry.get("business_purpose", ""))
        ws2.cell(row=row_i, column=3, value=entry["record_count"])
        ws2.cell(row=row_i, column=4, value=", ".join(entry.get("pii_columns", [])))
        ws2.cell(row=row_i, column=5, value=", ".join(entry.get("data_categories", [])))
        ws2.cell(row=row_i, column=6, value=entry.get("via") or "Primary table")

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    filename = f"dsar_{req.subject_id}_{req.created_at.strftime('%Y%m%d')}.xlsx"
    response = HttpResponse(
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


@api_view(["POST"])
def cancel_access_request(request, pk):
    req = get_object_or_404(AccessRequest, pk=pk)
    if req.status not in ("draft", "discovered", "report_ready"):
        return Response({"detail": "Cannot cancel a request in this state."}, status=400)
    req.status = "cancelled"
    req.save(update_fields=["status"])
    return Response({"detail": "Request cancelled."})
